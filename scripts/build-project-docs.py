#!/usr/bin/env python3
"""Build the project dossier from the reviewed Markdown sources (no app data)."""
from pathlib import Path
import html
import re
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, LongTable, TableStyle, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
TITLE = 'RegistroMedia — Documentazione completa e WBS'
DATE = '30 ottobre 2026'
NAVY = '#24304A'
ACCENT = '#3F57CD'
MUTED = '#626E80'
FONTROOT = Path('/usr/share/fonts/truetype/dejavu')
for name, file in [('DocSans', 'DejaVuSans.ttf'), ('DocSans-Bold', 'DejaVuSans-Bold.ttf'), ('DocSans-Oblique', 'DejaVuSans-Oblique.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(FONTROOT / file)))
pdfmetrics.registerFontFamily('DocSans', normal='DocSans', bold='DocSans-Bold', italic='DocSans-Oblique', boldItalic='DocSans-Bold')


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
            continue
        h = re.match(r'^(#{1,3})\s+(.*)$', line)
        if h:
            yield 'h' + str(len(h[1])), h[2]
            i += 1
            continue
        if line.startswith('- '):
            yield 'bullet', line[2:]
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||- )', lines[i].strip()):
            paragraph.append(lines[i].strip())
            i += 1
        yield 'p', ' '.join(paragraph)


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('**', '').replace('`', '')


def inline(text):
    text = escape(text)
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text


styles = {
    'p': ParagraphStyle('Body', fontName='DocSans', fontSize=9.1, leading=14, spaceAfter=8, textColor=colors.HexColor(NAVY)),
    'h1': ParagraphStyle('H1', fontName='DocSans-Bold', fontSize=23, leading=29, spaceAfter=17, spaceBefore=8, keepWithNext=True, textColor=colors.HexColor(NAVY)),
    'h2': ParagraphStyle('H2', fontName='DocSans-Bold', fontSize=15, leading=21, spaceAfter=11, spaceBefore=20, keepWithNext=True, textColor=colors.HexColor(NAVY)),
    'h3': ParagraphStyle('H3', fontName='DocSans-Bold', fontSize=11, leading=16, spaceAfter=8, spaceBefore=13, keepWithNext=True, textColor=colors.HexColor(ACCENT)),
    'bullet': ParagraphStyle('Bullet', fontName='DocSans', fontSize=9.1, leading=14, leftIndent=11, firstLineIndent=-9, spaceAfter=6, textColor=colors.HexColor(NAVY)),
    'cell': ParagraphStyle('Cell', fontName='DocSans', fontSize=8, leading=11.5, spaceAfter=0, textColor=colors.HexColor(NAVY)),
    'th': ParagraphStyle('TH', fontName='DocSans-Bold', fontSize=8, leading=11.5, textColor=colors.white),
    'role': ParagraphStyle('Role', fontName='DocSans', fontSize=9.1, leading=14, spaceAfter=8, keepWithNext=True, textColor=colors.HexColor(NAVY)),
    'cover_label': ParagraphStyle('CoverLabel', fontName='DocSans-Bold', fontSize=11, leading=16, textColor=colors.HexColor(ACCENT)),
    'cover_title': ParagraphStyle('CoverTitle', fontName='DocSans-Bold', fontSize=35, leading=43, textColor=colors.HexColor(NAVY)),
    'cover_sub': ParagraphStyle('CoverSub', fontName='DocSans', fontSize=15, leading=23, textColor=colors.HexColor(MUTED)),
}


class PageCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.states = []
    def showPage(self):
        self.states.append(dict(self.__dict__))
        self._startPage()
    def save(self):
        total = len(self.states)
        for state in self.states:
            self.__dict__.update(state)
            width, height = A4
            self.setStrokeColor(colors.HexColor('#DCE2EC'))
            self.line(54, 39, width - 54, 39)
            self.setFillColor(colors.HexColor(MUTED))
            self.setFont('DocSans', 7)
            self.drawString(54, 26, 'RegistroMedia | Documento 1.0 | ' + DATE)
            self.drawRightString(width - 54, 26, f'{self._pageNumber} / {total}')
            if self._pageNumber > 1:
                self.setFont('DocSans-Bold', 7)
                self.drawString(54, height - 32, 'REGISTROMEDIA / RELAZIONE E WBS / SOFTWARE 0.4.0')
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)


word = Document()
sec = word.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.9)
sec.left_margin = sec.right_margin = Cm(1.9)
normal = word.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15
for name, size in [('Title', 34), ('Heading 1', 23), ('Heading 2', 16), ('Heading 3', 12)]:
    style = word.styles[name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(NAVY[1:])
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(8)
word.core_properties.title = TITLE
word.core_properties.subject = 'Registro personale offline — architettura, manuale, verifiche e WBS'
word.core_properties.author = 'RegistroMedia'
header = sec.header.paragraphs[0]
run = header.add_run('REGISTROMEDIA  /  RELAZIONE E WBS  /  SOFTWARE 0.4.0')
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string(MUTED[1:])
footer = sec.footer.paragraphs[0]
footer.add_run('Documento 1.0 | ' + DATE + '    ·    Pagina ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(8)

story = [Spacer(1, 55), Paragraph('REGISTROMEDIA', styles['cover_label']), Spacer(1, 25),
         Paragraph('Documentazione<br/>completa del progetto', styles['cover_title']), Spacer(1, 25),
         Paragraph('Relazione tecnica, guida funzionale<br/>e Work Breakdown Structure', styles['cover_sub']),
         Spacer(1, 45), Paragraph('Software 0.4.0 · Prerelease personale offline', styles['p']),
         Paragraph('Windows · Android · Frontend web', styles['p']), Spacer(1, 40),
         Paragraph('Documento 1.0 — ' + DATE, styles['p']),
         Paragraph('Base: codice f17d702. Stati realizzati, verifiche e attività future esplicitamente distinti.', styles['p']),
         Spacer(1, 30), Paragraph('Nessun dato personale reale, password o chiave privata è incluso in questo documento.', styles['p']), PageBreak()]
word.add_paragraph('REGISTROMEDIA', 'Subtitle')
word.add_paragraph('Documentazione completa del progetto', 'Title')
word.add_paragraph('Relazione tecnica, guida funzionale e Work Breakdown Structure', 'Subtitle')
word.add_paragraph('Software 0.4.0 · Prerelease personale offline\nWindows · Android · Frontend web')
word.add_paragraph('Documento 1.0 — ' + DATE)
word.add_paragraph('Base: codice f17d702. Stati realizzati, verifiche e attività future esplicitamente distinti.')
word.add_paragraph('Nessun dato personale reale, password o chiave privata è incluso in questo documento.')
word.add_page_break()

width = A4[0] - 108
for source in ['PROJECT-DOCUMENTATION.md', 'WBS.md']:
    if source == 'WBS.md':
        story.append(PageBreak()); word.add_page_break()
    for kind, value in blocks((ROOT / 'docs' / source).read_text()):
        if kind == 'table':
            cols = len(value[0])
            ratios = {2: [.29, .71], 3: [.16, .69, .15] if source == 'WBS.md' else [.23, .32, .45], 4: [.13, .28, .39, .20]}.get(cols, [1 / cols] * cols)
            table_data = [[Paragraph(inline(cell), styles['th'] if ri == 0 else styles['cell']) for cell in row] for ri, row in enumerate(value)]
            table = LongTable(table_data, colWidths=[width * n for n in ratios], repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(NAVY)), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F5F9')]), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8), ('LINEBELOW', (0, 0), (-1, 0), .6, colors.HexColor(NAVY)), ('LINEBELOW', (0, 1), (-1, -1), .25, colors.HexColor('#DDE3EC'))]))
            story.extend([table, Spacer(1, 12)])
            t = word.add_table(rows=1, cols=cols)
            t.style = 'Table Grid'; t.autofit = False
            for col, ratio in zip(t.columns, ratios): col.width = Cm(17.2 * ratio)
            trpr = t.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            for ri, row in enumerate(value):
                cells = t.rows[0].cells if ri == 0 else t.add_row().cells
                row_xml = cells[0]._tc.getparent().get_or_add_trPr()
                no_split = OxmlElement('w:cantSplit'); row_xml.append(no_split)
                for ci, cell_text in enumerate(row):
                    cells[ci].width = Cm(17.2 * ratios[ci]); cells[ci].text = plain(cell_text)
                    for p in cells[ci].paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        for r in p.runs:
                            r.font.size = Pt(9)
                            if ri == 0: r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
                    if ri == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), NAVY[1:]); cells[ci]._tc.get_or_add_tcPr().append(shade)
            word.add_paragraph()
        else:
            content = ('• ' if kind == 'bullet' else '') + value
            story.append(Paragraph(inline(content), styles['role'] if content.startswith('Ruoli proposti:') else styles[kind]))
            if kind.startswith('h'):
                word.add_paragraph(plain(value), 'Heading ' + kind[1])
            else:
                p = word.add_paragraph(plain(value), 'List Bullet' if kind == 'bullet' else 'Normal')
                if value.startswith('Ruoli proposti:'): p.paragraph_format.keep_with_next = True

pdf = OUT / 'RegistroMedia-Documentazione-Completa.pdf'
docx = OUT / 'RegistroMedia-Documentazione-Completa.docx'
SimpleDocTemplate(str(pdf), pagesize=A4, leftMargin=54, rightMargin=54, topMargin=56, bottomMargin=54, title=TITLE, author='RegistroMedia').build(story, canvasmaker=PageCanvas)
word.save(docx)
print('Creati:', pdf, docx)
