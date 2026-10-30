#!/usr/bin/env python3
"""Render the retrospective, its real screenshots, and the animated presentation."""
from pathlib import Path
import re,html,json
from xml.sax.saxutils import escape
from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak,Image,LongTable,TableStyle,KeepTogether
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'docs/evolution';SHOTS=OUT/'screenshots'
for name,file in [('Votiva','DejaVuSans.ttf'),('VotivaBold','DejaVuSans-Bold.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
pdfmetrics.registerFontFamily('Votiva',normal='Votiva',bold='VotivaBold',italic='Votiva',boldItalic='VotivaBold')
INK='#193B31';GREEN='#14594E';MUTED='#64736A'
styles={
'p':ParagraphStyle('P',fontName='Votiva',fontSize=9.3,leading=14,spaceAfter=9,textColor=colors.HexColor(INK)),
'h1':ParagraphStyle('H1',fontName='VotivaBold',fontSize=28,leading=34,spaceAfter=18,textColor=colors.HexColor(INK)),
'h2':ParagraphStyle('H2',fontName='VotivaBold',fontSize=16,leading=22,spaceBefore=19,spaceAfter=11,keepWithNext=True,textColor=colors.HexColor(GREEN)),
'h3':ParagraphStyle('H3',fontName='VotivaBold',fontSize=11,leading=16,spaceBefore=12,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor(GREEN)),
'caption':ParagraphStyle('Caption',fontName='Votiva',fontSize=8,leading=11,spaceBefore=6,spaceAfter=14,textColor=colors.HexColor(MUTED)),
'cell':ParagraphStyle('Cell',fontName='Votiva',fontSize=8,leading=11,textColor=colors.HexColor(INK)),
'head':ParagraphStyle('Head',fontName='VotivaBold',fontSize=8,leading=11,textColor=colors.white),
'cover':ParagraphStyle('Cover',fontName='VotivaBold',fontSize=36,leading=43,textColor=colors.HexColor(INK),spaceAfter=20),
}
class Numbered(Canvas):
 def __init__(self,*a,**kw):super().__init__(*a,**kw);self.saved=[]
 def showPage(self):self.saved.append(dict(self.__dict__));self._startPage()
 def save(self):
  total=len(self.saved)
  for state in self.saved:
   self.__dict__.update(state);self.setFont('Votiva',8);self.setFillColor(colors.HexColor(MUTED));self.drawString(48,25,'VOTIVA / STORIA DEL PROGETTO / 30 OTTOBRE 2026');self.drawRightString(A4[0]-48,25,f'{self._pageNumber} / {total}');self.setStrokeColor(colors.HexColor('#dfe5dc'));self.line(48,A4[1]-34,A4[0]-48,A4[1]-34);Canvas.showPage(self)
  Canvas.save(self)
word=Document();sec=word.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
normal=word.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
for name in ['Heading 1','Heading 2','Heading 3']:word.styles[name].font.color.rgb=RGBColor.from_string('14594E')
word.core_properties.title='Da RegistroMedia a Votiva — storia illustrata del progetto';word.core_properties.author='Votiva'
sec.header.paragraphs[0].text='VOTIVA  /  EVOLUZIONE DEL PROGETTO'
f=sec.footer.paragraphs[0];f.add_run('30 ottobre 2026 · Pagina ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
WIDTH=A4[0]-96
story=[]
def para(text,kind='p'):
 story.append(Paragraph(escape(text),styles[kind]));word.add_paragraph(text,{'h1':'Heading 1','h2':'Heading 2','h3':'Heading 3'}.get(kind,'Normal'))
def figure(file,caption,maxheight=280):
 path=SHOTS/file
 if not path.exists():raise FileNotFoundError(path)
 w,h=PILImage.open(path).size;scale=min(WIDTH/w,maxheight/h)
 image=Image(str(path),width=w*scale,height=h*scale);story.append(KeepTogether([Spacer(1,8),image,Paragraph(escape(caption),styles['caption'])]))
 p=word.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(w*scale/72));word.add_paragraph(caption,'Caption')
para('DA REGISTROMEDIA A VOTIVA','h3');story.append(Spacer(1,22));story.append(Paragraph('Un’idea che cambia.<br/>Un prodotto che trova<br/>il proprio scopo.',styles['cover']));word.add_paragraph('Un’idea che cambia. Un prodotto che trova il proprio scopo.','Title')
para('Storia illustrata, cambiamenti, miglioramenti e compromessi.');para('Dall’idea Argo alla versione personale offline 1.0.1 · 30 ottobre 2026.');figure('03-votiva-101-desktop.png','Votiva 1.0.1 — interfaccia reale nel browser, con dati dimostrativi.',260)
story.append(PageBreak());word.add_page_break()
figure_map={
'4. Prima realizzazione':[('01-registromedia-03-accesso.png','Figura 1. Accesso di RegistroMedia 0.3, ricostruito dal commit bedd5f8.'),('01-registromedia-03-desktop.png','Figura 2. RegistroMedia 0.3: barra laterale, riepilogo e strumenti più numerosi.')],
'5. Prima semplificazione':[('02-registromedia-04-desktop.png','Figura 3. RegistroMedia 0.4: interfaccia semplificata, due schede e importazione.'),('02-registromedia-04-importazione.png','Figura 4. Il dialogo OCR della 0.4. Funzione esistente allora, rimossa dalla versione finale.')],
'7. La scelta finale':[('03-votiva-101-desktop.png','Figura 5. Votiva 1.0.1: nuovo nome, logo, materie e grafico.'),('04-votiva-voti-matematica.png','Figura 6. Voti della materia selezionata: esempio Matematica.'),('05-votiva-inserimento-voto.png','Figura 7. Inserimento voto: selettore con incrementi di 0,25.'),('06-votiva-periodi.png','Figura 8. Date dei periodi configurabili, ricorrenti ogni anno.')],
'8. Tutti i principali':[('07-votiva-scuro.png','Figura 9. Tema scuro: la struttura resta coerente con il tema chiaro.')],
'12. Sicurezza':[('11-votiva-account.png','Figura 10. Account locale, backup e gestione dei dati nella versione finale.'),('12-votiva-password.png','Figura 11. Mostra durante la digitazione e cambio password; nessuna password reale esposta.')],
'14. Verifiche':[('08-votiva-mobile.png','Figura 12. Vista mobile simulata a 390 × 844 pixel; non è una foto di un iPhone.'),('10-votiva-tablet.png','Figura 13. Vista tablet simulata a 768 × 1024 pixel.')],
}
def emit_figures(section):
 for key,figs in figure_map.items():
  if section.startswith(key):
   for file,caption in figs:figure(file,caption,310 if 'mobile' not in file else 340)
lines=(OUT/'STORIA-DEL-PROGETTO.md').read_text().splitlines();i=0;section=''
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r':?-+:?',x) for x in cells):rows.append(cells)
   i+=1
  n=len(rows[0]);ratios={3:[.23,.33,.44],4:[.26,.23,.23,.28]}.get(n,[1/n]*n)
  pdfrows=[[Paragraph(escape(c),styles['head' if j==0 else 'cell']) for c in row] for j,row in enumerate(rows)]
  table=LongTable(pdfrows,colWidths=[WIDTH*x for x in ratios],repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(GREEN)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#f1f5ef'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.25,colors.HexColor('#dfe5dc'))]));story.extend([table,Spacer(1,12)])
  t=word.add_table(rows=1,cols=n);t.style='Light Shading Accent 1'
  for c,v in zip(t.rows[0].cells,rows[0]):c.text=v
  props=t.rows[0]._tr.get_or_add_trPr();el=OxmlElement('w:tblHeader');props.append(el)
  for row in rows[1:]:
   for c,v in zip(t.add_row().cells,row):c.text=v
  continue
 h=re.match(r'^(#{1,3}) (.*)',line)
 if h:
  level=len(h[1]);title=h[2]
  if level==2:
   emit_figures(section);section=title
  para(title,'h'+str(level));i+=1;continue
 parts=[line];i+=1
 while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith(('#','|')):parts.append(lines[i].strip());i+=1
 para(' '.join(parts))
emit_figures(section)
para('Appendice: indice dei materiali','h2')
for text in ['STORIA-DEL-PROGETTO.md: fonte testuale modificabile.','Votiva-Storia-del-Progetto.pdf e .docx: relazione illustrata.','media/Votiva-Evoluzione-e-Demo.mp4: video con didascalie, senza audio.','media/Votiva-Anteprima.gif: breve estratto animato.','presentazione.html: presentazione animata con immagini, offline e navigabile.','screenshots/: 18 PNG reali di interfacce storiche e correnti.','capture-manifest.json e media/video-info.json: riferimenti delle catture e proprietà del video.'] :para(text)
SimpleDocTemplate(str(OUT/'Votiva-Storia-del-Progetto.pdf'),pagesize=A4,leftMargin=48,rightMargin=48,topMargin=49,bottomMargin=45,title='Da RegistroMedia a Votiva',author='Votiva').build(story,canvasmaker=Numbered)
word.save(OUT/'Votiva-Storia-del-Progetto.docx')
slides=[
('L’idea iniziale','Automatizzare le medie con i dati Argo','Le API non sono mai state attivate. Il collegamento al login non equivaleva a una sincronizzazione.','01-registromedia-03-accesso.png'),
('Il primo prodotto','RegistroMedia 0.3: più strumenti','Il registro personale comprendeva agenda, obiettivi e simulatore. Una gamma più ampia, con maggiore densità dell’interfaccia.','01-registromedia-03-desktop.png'),
('La prima semplificazione','RegistroMedia 0.4: due schede e temi','La schermata si semplifica. Agenda e simulatore diventano strumenti secondari.','02-registromedia-04-desktop.png'),
('Un esperimento reale','Importazione locale da screenshot','L’OCR prevedeva revisione e conferma. Aiutava con i voti storici, ma aggiungeva peso e complessità.','02-registromedia-04-importazione.png'),
('La scelta finale','Votiva 1.0.1: un compito più chiaro','Rimossi Argo, OCR, agenda e simulatore. Restano voti, medie e strumenti direttamente collegati.','03-votiva-101-desktop.png'),
('La registrazione','Voti ogni 0,25','Da 0,25 a 10. Tipo, peso, data e nota restano modificabili.','05-votiva-inserimento-voto.png'),
('La consultazione','Tutti i voti della materia','Tocca il nome di una materia, filtra il periodo e modifica i risultati.','04-votiva-voti-matematica.png'),
('Il calendario','Confini personalizzabili','Trimestre e pentamestre si adattano al profilo. Le impostazioni si ripetono ogni anno.','06-votiva-periodi.png'),
('Il design','Chiaro, scuro e layout responsive','Otto larghezze controllate automaticamente. Le viste mobile e tablet sono simulate nel browser.','07-votiva-scuro.png'),
('I dispositivi','Una vista adatta al telefono','Su iPhone occorre il sito HTTPS pubblicato, poi Safari → Condividi → Aggiungi alla schermata Home.','08-votiva-mobile.png'),
('Il bilancio','Più leggera, meno automatizzata','APK: 15,6 → 3,04 MB (circa −80,5%). Si perde l’OCR e si inseriscono i voti a mano. Nessuna misura di velocità attribuita.','10-votiva-tablet.png'),
('Il punto di arrivo','Un progresso con limiti dichiarati','Build e test automatici riusciti. Collaudo fisico, Pages e firme di produzione restano da completare.','03-votiva-101-desktop.png'),
]
css='''*{box-sizing:border-box}body{margin:0;font:16px system-ui;background:#112b24;color:#edf6e7}header{padding:24px 5vw;display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid #ffffff22}header strong{font-size:23px}header span{font-size:13px;color:#b6c9b8}.slide{display:none;min-height:calc(100dvh - 170px);padding:38px 5vw;grid-template-columns:minmax(240px,.8fr) minmax(0,1.4fr);align-items:center;gap:4vw}.slide.active{display:grid;animation:enter .55s ease both}.kicker{font-size:12px;color:#c5ed85;letter-spacing:3px;text-transform:uppercase}h1{font-size:clamp(30px,4vw,58px);line-height:1.1;letter-spacing:-1.5px;margin:22px 0}p{font-size:18px;line-height:1.65;color:#c2d3c3}.visual{background:#ffffff0c;border:1px solid #ffffff25;border-radius:16px;padding:12px;text-align:center;overflow:hidden}.visual img{max-width:100%;max-height:65vh;object-fit:contain;display:block;margin:auto;border-radius:8px}.caption{font-size:11px;color:#b9ccbb;margin-top:9px}footer{display:flex;justify-content:space-between;align-items:center;gap:16px;padding:18px 5vw;border-top:1px solid #ffffff22}button,a{font:inherit;color:inherit}button{border:1px solid #ffffff35;border-radius:9px;background:#ffffff0d;padding:11px 15px;cursor:pointer}button:hover,button:focus-visible{background:#ffffff25}nav{display:flex;gap:10px}#progress{height:3px;background:#c5ed85;transition:width .4s}.notes{font-size:11px;max-width:420px;color:#b9ccbb}a{color:#c5ed85}@keyframes enter{from{opacity:0;transform:translateY(13px)}to{opacity:1;transform:translateY(0)}}@media(max-width:750px){header{gap:20px}.slide.active{grid-template-columns:1fr;gap:24px;padding:25px 6vw}.visual img{max-height:52vh}h1{font-size:34px}p{font-size:16px}footer{flex-wrap:wrap}.notes{max-width:none}}@media(prefers-reduced-motion:reduce){.slide.active{animation:none}#progress{transition:none}}'''
parts=['<!doctype html><html lang="it"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Da RegistroMedia a Votiva — presentazione</title><style>'+css+'</style><header><strong>Votiva / Il percorso</strong><span>30 ottobre 2026 · Dati dimostrativi</span></header><div id="progress"></div><main>']
for i,(k,title,body,shot) in enumerate(slides):parts.append(f'<section class="slide{" active" if i==0 else ""}" aria-label="Diapositiva {i+1}"><div><div class="kicker">{html.escape(k)}</div><h1>{html.escape(title)}</h1><p>{html.escape(body)}</p></div><div class="visual"><img src="screenshots/{shot}" alt="{html.escape(title)}"><div class="caption">Interfaccia reale · screenshot ricostruito dai sorgenti · account e voti di esempio</div></div></section>')
parts.append('''</main><footer><nav><button id="prev" aria-label="Diapositiva precedente">←</button><button id="play">Avvia animazione</button><button id="next" aria-label="Diapositiva successiva">→</button></nav><span id="counter" aria-live="polite"></span><div class="notes">Frecce: naviga · Spazio: pausa/avvia. <a href="media/Votiva-Evoluzione-e-Demo.mp4">Apri il video</a>. PDF e Word contengono immagini statiche.</div></footer><script>
const slides=[...document.querySelectorAll('.slide')];let current=0,timer=null;const play=document.getElementById('play');function show(n){current=(n+slides.length)%slides.length;slides.forEach((s,i)=>{s.classList.toggle('active',i===current);s.setAttribute('aria-hidden',String(i!==current));});document.getElementById('counter').textContent=`${current+1} / ${slides.length}`;document.getElementById('progress').style.width=((current+1)/slides.length*100)+'%';}function toggle(){if(timer){clearInterval(timer);timer=null;play.textContent='Avvia animazione';}else{timer=setInterval(()=>show(current+1),8000);play.textContent='Pausa';}}document.getElementById('next').onclick=()=>show(current+1);document.getElementById('prev').onclick=()=>show(current-1);play.onclick=toggle;document.addEventListener('keydown',e=>{if(e.target.closest('button,a,input,select,textarea'))return;if(e.key==='ArrowRight')show(current+1);if(e.key==='ArrowLeft')show(current-1);if(e.code==='Space'){e.preventDefault();toggle();}});show(0);
</script></html>''')
(OUT/'presentazione.html').write_text(''.join(parts))
print('Generated PDF, DOCX and HTML presentation')
